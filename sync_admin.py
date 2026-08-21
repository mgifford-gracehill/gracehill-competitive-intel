"""Apply human edits from Grace-Hill-CI-Admin.xlsx back into the data layer.

Human-owned fields always win. Machine-owned fields (facts, pricing, strengths,
gaps, competitor feature values, news, sources) are never touched here.

Usage:  python3 sync_admin.py [path-to-xlsx]
"""
import json, sys, pathlib, re, os
import flatten as FL
from openpyxl import load_workbook

# The Drive connector renders the workbook as markdown and collapses newlines inside a cell,
# so a multi-line list arrives as one space-joined string. Applying that would silently destroy
# the list. In that mode we compare on normalised text and refuse to write list fields that
# genuinely differ, rather than writing a corrupted single-item list.
LOSSY = os.environ.get('SYNC_SOURCE') == 'drive'
def norm(x):
    if isinstance(x, list): x = ' '.join(x)
    return re.sub(r'\s+', ' ', str(x or '')).strip()

XLSX = sys.argv[1] if len(sys.argv) > 1 else 'Grace-Hill-CI-Admin.xlsx'
wb = load_workbook(XLSX, data_only=True)
changes, warnings = [], []
def lines(v): return [x.strip() for x in str(v or '').split('\n') if x.strip()]
def clean(v):
    v = None if v is None else str(v).strip()
    return v or None

# ---------------- Competitors ----------------
C = json.load(open('data/competitors.json'))
byid = {c['id']: c for c in C}
ws = wb['Competitors']
hdr = [c.value for c in ws[1]]
COL = {h: i for i, h in enumerate(hdr)}
def g(row, name): return row[COL[name]].value if name in COL else None
retired = []
for row in ws.iter_rows(min_row=2):
    cid = clean(g(row, 'id'))
    if not cid: continue
    c = byid.get(cid)
    if not c:
        warnings.append(f"Row for unknown id '{cid}' — ignored. Use the Requests tab to add a competitor."); continue
    if (clean(g(row, 'Status')) or 'Active').lower() == 'retire':
        retired.append(cid); continue
    def setf(path, val, label):
        if val is None: return
        tgt, key = c, path[-1]
        for p in path[:-1]: tgt = tgt.setdefault(p, {})
        if tgt.get(key) != val:
            tgt[key] = val; changes.append(f"{cid}: {label}")
    setf(['name'], clean(g(row,'Name')), 'name')
    setf(['toolCategory'], clean(g(row,'What it is (tool category)')), 'tool category')
    setf(['ciCategory'], clean(g(row,'Competitive category')), 'competitive category')
    setf(['tier'], clean(g(row,'Tier')), 'tier')
    setf(['verdict','stance'], clean(g(row,'Stance')), 'stance')
    setf(['verdict','line'], clean(g(row,'Verdict line (one sentence a rep reads in 5 seconds)')), 'verdict line')
    setf(['verdict','doInstead'], clean(g(row,'Do instead (required if Do not compete / Partner)')), 'do instead')
    before = len(changes)
    for key, col, label in [('why',"Why this is / isn't a fight",'why'),('say','Say this','say'),
                            ('avoid','Avoid','avoid'),('redirect','Then redirect','redirect')]:
        setf(['whenItComesUp',key], clean(g(row,col)), f'when-it-comes-up · {label}')
    # only promote off 'derived' when a human actually rewrote the guidance
    if len(changes) > before:
        w = c.get('whenItComesUp') or {}
        if w.pop('derived', None): changes.append(f"{cid}: guidance now hand-written")
    for key, col in [('whyWeWin','Win with (one per line)'),('askThis','Ask this (one per line)'),('doNotSay','Do not say (one per line)')]:
        raw = g(row,col); v = lines(raw)
        if not v or v == c.get(key): continue
        if LOSSY:
            if norm(raw) == norm(c.get(key)): continue          # unchanged, just flattened in transit
            warnings.append(f"{cid}: '{col.split(' (')[0]}' changed but the Drive view flattens line breaks — export the sheet as .xlsx and sync that so the list survives. Not applied.")
            continue
        c[key] = v; changes.append(f"{cid}: {col.split(' (')[0]}")
    raw_ov = g(row,'Overlaps (Product:Degree, one per line)')
    ov = []
    for ln in lines(raw_ov):
        if ':' in ln:
            p, d = ln.split(':', 1); ov.append({'product': p.strip(), 'degree': d.strip().title()})
        else: warnings.append(f"{cid}: overlap '{ln}' ignored — use Product:Degree")
    cur_ov = [f"{o['product']}:{o['degree']}" for o in (c.get('overlaps') or [])]
    if ov and ov != c.get('overlaps'):
        if LOSSY and norm(raw_ov) != norm(cur_ov):
            warnings.append(f"{cid}: overlaps changed but the Drive view flattens line breaks — export as .xlsx to apply. Not applied.")
        elif LOSSY:
            pass
        else:
            c['overlaps'] = ov; changes.append(f"{cid}: overlaps")
    owner = clean(g(row,'Owner'))
    if owner: c.setdefault('health', {})['owner'] = owner
    note = clean(g(row,'Notes for Claude'))
    if note: warnings.append(f"NOTE FOR CLAUDE · {cid}: {note}")
    # integrity: a non-compete stance must tell the rep what to do instead
    v = c.get('verdict') or {}
    if v.get('stance') in ('Do not compete','Partner') and not v.get('doInstead'):
        warnings.append(f"{cid}: stance is '{v['stance']}' but 'Do instead' is empty — a rep is left with nothing to say.")
if retired:
    C = [c for c in C if c['id'] not in retired]
    changes.append(f"retired from the app: {', '.join(retired)}")
    pathlib.Path('data/retired.json').write_text(json.dumps([byid[r] for r in retired], indent=2, ensure_ascii=False))
json.dump(C, open('data/competitors.json','w'), indent=2, ensure_ascii=False)

# ---------------- Our Capabilities (the Grace Hill column) ----------------
GHB = json.load(open('data/gh-baseline.json'))
ws = wb['Our Capabilities']; hdr = [c.value for c in ws[1]]; COL = {h:i for i,h in enumerate(hdr)}
for row in ws.iter_rows(min_row=2):
    fid = clean(g(row,'Feature id'))
    if not fid or fid not in GHB: continue
    for key, col in [('support','Grace Hill support'),('note','What we actually do (rep-sayable)'),('proof','Proof'),('source','Source')]:
        v = clean(g(row,col))
        if v is not None and GHB[fid].get(key) != v:
            GHB[fid][key] = v; changes.append(f"our capability {fid}: {key}")
json.dump(GHB, open('data/gh-baseline.json','w'), indent=2, ensure_ascii=False)

# ---------------- Grace Hill products & capabilities ----------------
GH = json.load(open('data/gh-index.json')); gid = {x['id']: x for x in GH}
ws = wb['Grace Hill']; hdr = [c.value for c in ws[1]]; COL = {h:i for i,h in enumerate(hdr)}
for row in ws.iter_rows(min_row=2):
    i = clean(g(row,'id'))
    if not i or i not in gid: continue
    x = gid[i]
    for key, col in [('name','Name'),('parent','Part of'),('line','Product line'),
                     ('whatItIs','What it is'),('why','Why it matters in a deal'),('note','Be precise about this')]:
        v = clean(g(row,col))
        if v is not None and x.get(key) != v: x[key] = v; changes.append(f"{i}: {key}")
    for key, col in [('aka','Also called (one per line)'),('competitors','Competitors it answers (one per line)')]:
        raw = g(row,col); v = lines(raw)
        if not v or v == x.get(key): continue
        if LOSSY:
            if norm(raw) != norm(x.get(key)):
                warnings.append(f"{i}: '{key}' changed but the Drive view flattens line breaks — export as .xlsx to apply. Not applied.")
            continue
        x[key] = v; changes.append(f"{i}: {key}")
json.dump(GH, open('data/gh-index.json','w'), indent=2, ensure_ascii=False)

# ---------------- Content tabs: pitch, SPICED, plays, framework, personas, ICP ----------------
CONTENT = [('Pitch','data/pitch.json'),('SPICED','data/spiced.json'),('Plays','data/plays.json'),
           ('Framework','data/framework.json'),('Personas','data/personas.json'),('ICP','data/icp.json')]
for tab, path in CONTENT:
    if tab not in wb.sheetnames: continue
    doc = json.load(open(path))
    ws = wb[tab]
    hrow = None
    for r in range(1, 8):
        if str(ws.cell(r,1).value or '').strip() == 'key': hrow = r; break
    if hrow is None:
        warnings.append(f"{tab}: no header row found — skipped"); continue
    n = 0
    for r in range(hrow+1, ws.max_row+1):
        key = clean(ws.cell(r,1).value)
        if not key: continue
        val = ws.cell(r,4).value
        val = '' if val is None else str(val)
        cur = FL.get_path(doc, key)
        cur = '' if cur is None else cur
        if not isinstance(cur, (str, int, float)):
            warnings.append(f"{tab}: key '{key}' does not point at a text field — skipped"); continue
        if str(cur) != val:
            FL.set_path(doc, key, val); n += 1; changes.append(f"{tab}: {key}")
    if n: json.dump(doc, open(path,'w'), indent=2, ensure_ascii=False)

# ---------------- Requests ----------------
ws = wb['Requests']; hdr = [c.value for c in ws[3]]; COL = {h:i for i,h in enumerate(hdr)}
open_reqs = []
for row in ws.iter_rows(min_row=4):
    if (clean(g(row,'Status')) or '').lower() != 'open': continue
    if not clean(g(row,'Competitor / product')): continue
    open_reqs.append({k: clean(g(row,k)) for k in ['Date','Requested by','Type','Competitor / product','What you want','Source or URL (optional)']})

print(json.dumps({'changesApplied': len(changes), 'changes': changes[:40],
                  'openRequests': open_reqs, 'warnings': warnings}, indent=2, ensure_ascii=False))
