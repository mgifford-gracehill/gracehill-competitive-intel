"""Sync from the live Google Sheet.

The Drive connector renders the whole workbook as markdown tables with no tab
separators, so tabs are segmented by their unique header row. The parsed result is
written to a temporary .xlsx with the same headers the admin workbook uses, then
handed to sync_admin.py — so the Drive path and the file path share one tested
mapping instead of drifting apart.

Usage:  python3 sync_from_drive.py drive_dump.txt
"""
import re, sys, json, pathlib, subprocess
from openpyxl import Workbook

SRC = sys.argv[1] if len(sys.argv) > 1 else 'drive_dump.txt'
lines = pathlib.Path(SRC).read_text().split('\n')

def unesc(s):
    s = re.sub(r'\\([^\w\s])', r'\1', s or '')   # markdown escapes any punctuation
    return s.replace('<br>', '\n').strip()

def cells(line):
    if not line.startswith('|'): return None
    parts = re.split(r'(?<!\\)\|', line)          # split on unescaped pipes only
    parts = parts[1:-1] if line.rstrip().endswith('|') else parts[1:]
    return [unesc(p) for p in parts]

HEADERS = {
 'Competitors':      lambda c: c[:2] == ['id','Name'] and 'What it is (tool category)' in c,
 'Our Capabilities': lambda c: c[:1] == ['Feature id'],
 'Grace Hill':       lambda c: c[:3] == ['id','Name','Type'],
 'Requests':         lambda c: c[:2] == ['Date','Requested by'],
}
CONTENT_NOTES = [('Pitch','The 30/3/30 pitch'),('SPICED','Discovery questions, deal-review'),
                 ('Plays','The demand playbook'),('Framework','How we classify a competitor'),
                 ('Personas','The nine buyer personas'),('ICP','Ideal customer profile, fit')]

# locate each tab's header line
starts = {}
pending_content = None
for i, ln in enumerate(lines):
    c = cells(ln)
    if not c: continue
    for name, test in HEADERS.items():
        if name not in starts and test(c): starts[name] = i
    for name, note in CONTENT_NOTES:
        if name not in starts and c and note in (c[0] or ''): pending_content = name
    if pending_content and c[:4] == ['key','Section','Field','Value']:
        starts[pending_content] = i; pending_content = None

order = sorted(starts.items(), key=lambda kv: kv[1])
bounds = {}
for idx, (name, st) in enumerate(order):
    end = order[idx+1][1] - 1 if idx+1 < len(order) else len(lines)
    bounds[name] = (st, end)

tabs = {}
for name, (st, end) in bounds.items():
    hdr = cells(lines[st])
    rows = []
    for ln in lines[st+1:end]:
        c = cells(ln)
        if not c: continue
        if set(''.join(c).replace(':','').replace('-','')) <= {''}: continue   # md separator row
        rows.append((c + ['']*len(hdr))[:len(hdr)])
    tabs[name] = {'header': hdr, 'rows': rows}

# rebuild a workbook sync_admin.py can consume
wb = Workbook(); wb.remove(wb.active)
for name in ['Competitors','Our Capabilities','Grace Hill','Pitch','SPICED','Plays','Framework','Personas','ICP']:
    if name not in tabs: continue
    ws = wb.create_sheet(name)
    if name in ('Pitch','SPICED','Plays','Framework','Personas','ICP'):
        ws.append([]); ws.append([])          # sync looks for the 'key' header within the first rows
    ws.append(tabs[name]['header'])
    for r in tabs[name]['rows']: ws.append(r)
if 'Requests' in tabs:
    ws = wb.create_sheet('Requests'); ws.append([]); ws.append([])
    ws.append(tabs['Requests']['header'])
    for r in tabs['Requests']['rows']: ws.append(r)
wb.save('.drive-sync.xlsx')

print(json.dumps({'tabsParsed': {k: len(v['rows']) for k, v in tabs.items()},
                  'missing': [t for t in list(HEADERS)+[n for n,_ in CONTENT_NOTES] if t not in tabs]},
                 indent=2, ensure_ascii=False))
print('--- applying ---')
subprocess.run([sys.executable, 'sync_admin.py', '.drive-sync.xlsx'], env={**__import__('os').environ, 'SYNC_SOURCE': 'drive'})
