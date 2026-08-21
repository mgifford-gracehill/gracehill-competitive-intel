import json, pathlib
import flatten as FL
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

C=json.load(open('data/competitors.json'))
GH=json.load(open('data/gh-index.json'))
GHB=json.load(open('data/gh-baseline.json'))
FEAT={}
for l in json.load(open('data/features.json'))['lines']:
    for f in l['features']: FEAT[f['id']]={**f,'line':l['line']}

NAVY='FF262262'; RED='FFFF0000'; YELLOW='FFFFF7CC'; GREY='FFF0F0F3'; WHITE='FFFFFFFF'
hdr=PatternFill('solid',fgColor=NAVY); edit=PatternFill('solid',fgColor=YELLOW); ro=PatternFill('solid',fgColor=GREY)
HF=Font(name='Arial',size=10,bold=True,color=WHITE)
BF=Font(name='Arial',size=10); BOLD=Font(name='Arial',size=10,bold=True)
TOP=Alignment(vertical='top',wrap_text=True)
thin=Side(style='thin',color='FFD9D9DE'); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
def nl(a): return '\n'.join(a or [])
def style(ws,ncols,nrows,editable_cols):
    for c in range(1,ncols+1):
        cell=ws.cell(1,c); cell.fill=hdr; cell.font=HF; cell.alignment=Alignment(vertical='center',wrap_text=True)
    for r in range(2,nrows+2):
        for c in range(1,ncols+1):
            cell=ws.cell(r,c); cell.font=BF; cell.alignment=TOP; cell.border=BORD
            cell.fill = edit if c in editable_cols else ro
    ws.freeze_panes='C2'; ws.row_dimensions[1].height=30

wb=Workbook()

# ---------- README ----------
ws=wb.active; ws.title='README'
rows=[
 ('Grace Hill Competitive Intelligence — Admin Sheet',None),
 ('',None),
 ('This sheet is the human-owned layer of the competitive intelligence app.',None),
 ('Anything you type in a YELLOW cell wins. Automated research will never overwrite it.',None),
 ('GREY cells are keys or machine-maintained context — do not edit them.',None),
 ('',None),
 ('HOW IT WORKS',None),
 ('1. Edit a yellow cell, or add a row to the Requests tab.',None),
 ('2. Tell Claude "sync the admin sheet". Claude reads this file and applies your edits.',None),
 ('3. The app rebuilds and the change is live in the next published version.',None),
 ('',None),
 ('WHAT IS HUMAN-OWNED (yours)',None),
 ('Stance, tier, category, verdict line, when-it-comes-up guidance, win-with, ask-this, do-not-say,',None),
 ('overlaps, owner, status, and the Grace Hill capability scores on the Our Capabilities tab.',None),
 ('',None),
 ('WHAT IS MACHINE-OWNED (research maintains it)',None),
 ('Company facts, funding, ownership, pricing, their published strengths and gaps, competitor',None),
 ('feature values, news, and sources. These are re-verified on a schedule and are not in this sheet.',None),
 ('',None),
 ('MULTI-VALUE CELLS',None),
 ('Where a cell holds a list, put one item per line (Alt+Enter). Overlaps use Product:Degree,',None),
 ('one per line — for example  Training:High  then  Policies:Partial',None),
 ('',None),
 ('TO ADD A COMPETITOR',None),
 ('Use the Requests tab. Give the name and a URL if you have one; Claude researches the rest.',None),
 ('',None),
 ('TO RETIRE A COMPETITOR',None),
 ('Set Status to Retire on the Competitors tab. It leaves the app but the record is kept.',None),
 ('',None),
 ('THE CONTENT TABS',None),
 ('Pitch, SPICED, Plays, Framework, Personas and ICP each use the same layout:',None),
 ('one row per field, with the editable text in the Value column. Edit the Value, leave key alone.',None),
 ('The key column is how the sync finds the field — changing it breaks the link.',None),
 ('Numbered rows (…[0], …[1]) are list items. Add a row with the next number to add an item.',None),
]
for i,(t,_) in enumerate(rows,1):
    c=ws.cell(i,1,t); c.font=BOLD if (i==1 or t.isupper() and t) else BF
    if i==1: c.font=Font(name='Arial',size=14,bold=True,color=NAVY)
ws.column_dimensions['A'].width=105
ws['C1']='COUNTS'; ws['C1'].font=BOLD
ws['C2']='Competitors'; ws['D2']='=COUNTA(Competitors!A2:A2000)'
ws['C3']='We compete'; ws['D3']='=COUNTIF(Competitors!F2:F2000,"Compete")'
ws['C4']='Do not compete'; ws['D4']='=COUNTIF(Competitors!F2:F2000,"Do not compete")'
ws['C5']='Marked Retire'; ws['D5']='=COUNTIF(Competitors!R2:R2000,"Retire")'
ws['C6']='Open requests'; ws['D6']='=COUNTIF(Requests!G3:G500,"Open")'
for r in range(2,7):
    ws.cell(r,3).font=BF; ws.cell(r,4).font=BF
ws.column_dimensions['C'].width=22; ws.column_dimensions['D'].width=10

# ---------- COMPETITORS ----------
ws=wb.create_sheet('Competitors')
COLS=['id','Name','What it is (tool category)','Competitive category','Tier','Stance','Verdict line (one sentence a rep reads in 5 seconds)',
 'Do instead (required if Do not compete / Partner)','Why this is / isn\'t a fight','Say this','Avoid','Then redirect',
 'Win with (one per line)','Ask this (one per line)','Do not say (one per line)','Overlaps (Product:Degree, one per line)',
 'Owner','Status','Last reviewed','Notes for Claude']
ws.append(COLS)
order={'Tier 1':0,'Tier 2':1,'Tier 3':2,'Watchlist':3}
for c in sorted(C,key=lambda x:(order.get(x.get('tier'),9),x['name'])):
    v=c.get('verdict') or {}; w=c.get('whenItComesUp') or {}
    ws.append([c['id'],c['name'],c.get('toolCategory'),c.get('ciCategory'),c.get('tier'),v.get('stance'),
      v.get('line'),v.get('doInstead'),w.get('why'),w.get('say'),w.get('avoid'),w.get('redirect'),
      nl(c.get('whyWeWin')),nl(c.get('askThis')),nl(c.get('doNotSay')),
      nl([f"{o['product']}:{o['degree']}" for o in (c.get('overlaps') or [])]),
      (c.get('health') or {}).get('owner'),'Active',(c.get('health') or {}).get('lastFullReview'),None])
EDIT=set(range(2,19))|{20}; EDIT.discard(19)
style(ws,len(COLS),len(C),EDIT)
widths=[22,26,34,18,11,15,52,44,44,52,40,40,52,44,44,26,16,11,14,34]
for i,wd in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=wd
for rng,vals in [('D2:D2000','"Direct,Indirect,Replacement,Potential,Status Quo,Partner-Competitor"'),
                 ('E2:E2000','"Tier 1,Tier 2,Tier 3,Watchlist"'),
                 ('F2:F2000','"Compete,Partial overlap,Do not compete,Partner"'),
                 ('R2:R2000','"Active,Retire"')]:
    dv=DataValidation(type='list',formula1=vals,allow_blank=True); ws.add_data_validation(dv); dv.add(rng)

# ---------- OUR CAPABILITIES ----------
ws=wb.create_sheet('Our Capabilities')
ws.append(['Feature id','Product line','Capability','Grace Hill support','What we actually do (rep-sayable)','Proof','Source','Verified on'])
n=0
for fid,g in GHB.items():
    f=FEAT.get(fid,{})
    ws.append([fid,f.get('line'),f.get('label'),g.get('support'),g.get('note'),g.get('proof'),g.get('source'),g.get('verifiedOn')]); n+=1
style(ws,8,n,{4,5,6})
for i,wd in enumerate([26,20,40,16,60,42,42,13],1): ws.column_dimensions[get_column_letter(i)].width=wd
dv=DataValidation(type='list',formula1='"Yes,Partial,No,Roadmap,Unknown"',allow_blank=True); ws.add_data_validation(dv); dv.add('D2:D200')

# ---------- GRACE HILL ----------
ws=wb.create_sheet('Grace Hill')
ws.append(['id','Name','Type','Part of','Product line','Also called (one per line)','What it is','Why it matters in a deal','Be precise about this','Competitors it answers (one per line)'])
for g in GH:
    ws.append([g['id'],g['name'],g.get('type'),g.get('parent'),g.get('line'),nl(g.get('aka')),g.get('whatItIs'),g.get('why'),g.get('note'),nl(g.get('competitors'))])
style(ws,10,len(GH),{2,4,5,6,7,8,9,10})
for i,wd in enumerate([26,28,13,26,20,24,52,60,52,34],1): ws.column_dimensions[get_column_letter(i)].width=wd

# ---------- REQUESTS ----------
ws=wb.create_sheet('Requests')
ws['A1']='Add a competitor, flag something wrong, or ask for a change. Claude reads this tab on every sync.'
ws['A1'].font=Font(name='Arial',size=10,bold=True,italic=True,color=NAVY)
ws.append([])
ws.append(['Date','Requested by','Type','Competitor / product','What you want','Source or URL (optional)','Status','Claude notes'])
ws.append(['2026-08-11','Mandy Gifford','Add competitor','Elevate Mystery Shopping','New multifamily shop firm founded by a former Director of Training. Build a record and tell us whether it is a real threat.','https://elevatemysteryshopping.com/','Open','Example row — delete or leave.'])
for c in range(1,9):
    cell=ws.cell(3,c); cell.fill=hdr; cell.font=HF; cell.alignment=Alignment(vertical='center',wrap_text=True)
for r in range(4,60):
    for c in range(1,9):
        cell=ws.cell(r,c); cell.font=BF; cell.alignment=TOP; cell.border=BORD
        cell.fill = ro if c==8 else edit
for i,wd in enumerate([13,20,20,30,64,34,12,40],1): ws.column_dimensions[get_column_letter(i)].width=wd
ws.freeze_panes='A4'; ws.row_dimensions[3].height=28
dv=DataValidation(type='list',formula1='"Add competitor,Fix something wrong,Change guidance,Retire competitor,Add capability,Other"',allow_blank=True)
ws.add_data_validation(dv); dv.add('C4:C200')
dv2=DataValidation(type='list',formula1='"Open,Done,Rejected"',allow_blank=True); ws.add_data_validation(dv2); dv2.add('G4:G200')


# ---------- CONTENT TABS (pitch, SPICED, plays, framework, personas, ICP) ----------
FW_KEYS = ('foundationalQuestion','categories','tiers','statusQuoForms','joltFinding','liveDealGuidance')
def content_tab(title, data, note, filter_top=None):
    rows = FL.flatten({k:v for k,v in data.items() if not filter_top or k in filter_top} if isinstance(data,dict) else data)
    rows = [r for r in rows if r[0].split('.')[0] not in ('asOf','source','status') or True]
    ws = wb.create_sheet(title)
    ws['A1'] = note; ws['A1'].font = Font(name='Arial',size=10,bold=True,italic=True,color=NAVY)
    ws.append([]) ; ws.append(['key','Section','Field','Value'])
    for k,sec,item,val in rows: ws.append([k,sec,item,val])
    for c in range(1,5):
        cell=ws.cell(3,c); cell.fill=hdr; cell.font=HF; cell.alignment=Alignment(vertical='center',wrap_text=True)
    for r in range(4,len(rows)+4):
        for c in range(1,5):
            cell=ws.cell(r,c); cell.font=BF; cell.alignment=TOP; cell.border=BORD
            cell.fill = edit if c==4 else ro
    for i,wd in enumerate([34,30,30,110],1): ws.column_dimensions[get_column_letter(i)].width=wd
    ws.freeze_panes='D4'; ws.row_dimensions[3].height=26
    return len(rows)

counts={}
counts['Pitch']      = content_tab('Pitch', json.load(open('data/pitch.json')),
  'The 30/3/30 pitch. Edit any Value cell — the wording here is what reps read in the app. Retired language lives under "Language to avoid".')
counts['SPICED']     = content_tab('SPICED', json.load(open('data/spiced.json')),
  'Discovery questions, deal-review checklist and the Stage 2 exit bar. "Ask it live" is what a rep says in the room; "Deal review" is what they must be able to answer afterward.')
counts['Plays']      = content_tab('Plays', json.load(open('data/plays.json')),
  'The demand playbook. Each play has a segment, a budget owner and a note. Marked WIP until the bundles launch.')
counts['Framework']  = content_tab('Framework', json.load(open('data/framework.json')),
  'How we classify a competitor: the six categories, the tiers, the four forms of status quo, and the live-deal cautions.', FW_KEYS)
counts['Personas']   = content_tab('Personas', json.load(open('data/personas.json')),
  'The nine buyer personas. "Cares about", the deck framing and "Competitors they will raise" all surface in the app.')
counts['ICP']        = content_tab('ICP', json.load(open('data/icp.json')),
  'Ideal customer profile, fit signals and the four problem-based segments.')
print('content tabs:', counts)

wb.save('Grace-Hill-CI-Admin.xlsx')
print('saved Grace-Hill-CI-Admin.xlsx')
print('tabs:', wb.sheetnames, '| competitors:',len(C),'| capabilities:',n,'| gh entries:',len(GH))
